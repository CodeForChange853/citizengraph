"""Splitter on the CHO (HEALTH) and CSWDO layouts: Office/Section, N/A, bullets, wrapped rows."""

from pathlib import Path

import openpyxl
import pytest

from citizengraph.parsing.splitter import ServiceDraft, split_workbook

RAW = Path(__file__).resolve().parents[1] / "data" / "raw"


@pytest.fixture(scope="module")
def cho() -> list[ServiceDraft]:
    return split_workbook(RAW / "CYPCC_HEALTH.xlsx")


@pytest.fixture(scope="module")
def cswdo() -> list[ServiceDraft]:
    return split_workbook(RAW / "CYPCC_SOCIALWELFARE.xlsx")


def _merge_row(ws, row: int) -> None:
    for a, b in (("C", "E"), ("F", "G"), ("H", "I"), ("J", "L")):
        ws.merge_cells(f"{a}{row}:{b}{row}")


def make_health_sheet(tmp_path: Path) -> Path:
    """Two services in the HEALTH layout: N/A checklist + bullets, and a real checklist."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HEALTH"
    ws["A1"] = "CITY HEALTH OFFICE"
    ws["A2"] = "1. Sample Care"
    ws["A3"], ws["D3"] = "Office/Section", "City Health Office"
    ws["A4"], ws["D4"] = "Classification", "Simple"
    ws["A5"], ws["D5"] = "Type of Transaction", "Government to Citizen"
    ws["A6"], ws["D6"] = "Who may avail", "A long free-text description of who may avail."
    ws["A7"], ws["F7"] = "CHECKLIST OF REQUIREMENTS", "WHERE TO SECURE"
    ws["A8"], ws["F8"] = "N/A", "N/A"
    ws["A9"], ws["C9"], ws["F9"], ws["H9"], ws["J9"] = (
        "CLIENT STEPS",
        "AGENCY ACTION",
        "FEES TO BE PAID",
        "PROCESSING TIME",
        "PERSON RESPONSIBLE",
    )
    ws["A10"], ws["C10"], ws["F10"], ws["H10"], ws["J10"] = (
        "Receiving",
        "• Client is taken in\n•\tRecord is retrieved and the vital signs\ntaken\n• Logged",
        "None ",
        "10 min",
        "Nurse",
    )
    ws["A11"], ws["C11"], ws["F11"], ws["H11"], ws["J11"] = (
        "Payment",
        "Client pays",
        "250",
        "2 mins",
        "Cashier",
    )
    for r in (10, 11):
        _merge_row(ws, r)
    ws["A12"], ws["H12"] = "TOTAL", "12 min"

    ws["A14"] = "2. Split Service"
    ws["A15"], ws["D15"] = "Office/Section", "City Health Office"
    ws["A16"], ws["D16"] = "Classification", "Simple"
    ws["A17"], ws["D17"] = "Type of Transaction", "Government to Citizen"
    ws["A18"], ws["D18"] = "Who may avail", "Anyone"
    ws["A19"], ws["F19"] = "CHECKLIST OF REQUIREMENTS", "WHERE TO SECURE"
    ws["A20"], ws["F20"] = "Request letter", "Philippine National Police"
    ws["A21"], ws["F21"] = "Application Form", "N/A"
    ws["A22"], ws["C22"], ws["F22"], ws["H22"], ws["J22"] = (
        "CLIENT STEPS",
        "AGENCY ACTION",
        "FEES TO BE PAID",
        "PROCESSING TIME",
        "PERSON RESPONSIBLE",
    )
    # a step whose label and person sit on one row and fee/time on the next
    ws["A23"], ws["C23"], ws["J23"] = "Receiving", "Client is taken in", "Nurse"
    ws["C24"], ws["F24"], ws["H24"] = "• Assessed\n• Advised", "None", "1.15 min"
    ws["A25"], ws["C25"], ws["F25"], ws["H25"], ws["J25"] = (
        "Consultation",
        "• Examined\n• Health history taken",
        "P30.00",
        "2-3 days",
        "Physician",
    )
    ws["C26"] = "Report made and signed."  # unmarked extra line, no fee/time/person
    ws["A27"], ws["F27"], ws["H27"] = "TOTAL", "₱30.00", "3 dsys"
    for r in (23, 24, 25, 26):
        _merge_row(ws, r)
    path = tmp_path / "CYPCC_SAMPLE.xlsx"
    wb.save(path)
    return path


def test_synthetic_office_section_layout(tmp_path):
    a, b = split_workbook(make_health_sheet(tmp_path))
    assert (a.draft_id, b.draft_id) == ("CYPCC-01", "CYPCC-02")
    assert a.title == "1. Sample Care" and a.name == "Sample Care"
    assert a.office == "City Health Office"
    assert a.classification == "SIMPLE"
    assert a.transaction_type == "Government to Citizen"
    assert a.who_may_avail == "A long free-text description of who may avail."


def test_na_checklist_is_no_requirements_not_a_requirement(tmp_path):
    a, b = split_workbook(make_health_sheet(tmp_path))
    assert a.requirements == []
    assert any(f["code"] == "no_requirements_listed" for f in a.flags)
    assert [r.text for r in b.requirements] == ["Request letter", "Application Form"]
    assert b.requirements[0].where_to_secure == "Philippine National Police"
    assert b.requirements[1].where_to_secure is None
    codes = [f["code"] for f in b.flags]
    assert "where_to_secure_na" in codes and "where_to_secure_missing" not in codes


def test_bullets_are_split_into_a_list_and_wrapped_lines_are_joined(tmp_path):
    a, _ = split_workbook(make_health_sheet(tmp_path))
    step = a.steps[0]
    assert step.agency_action_items == [
        "Client is taken in",
        "Record is retrieved and the vital signs taken",
        "Logged",
    ]
    assert step.agency_action.startswith("• Client is taken in")  # raw text is kept
    assert a.steps[1].agency_action_items is None  # no bullets, no list
    assert "agency_action_items" not in a.steps[1].to_dict()


def test_fee_and_time_cells_from_header_columns(tmp_path):
    a, _ = split_workbook(make_health_sheet(tmp_path))
    assert a.steps[0].fees.status == "none"
    assert a.steps[0].time.components[0].value_max == 10
    assert a.steps[1].fees.items[0].amount_max == 250
    assert a.steps[1].fees.items[0].note == "no currency sign in source"
    assert any(f["code"] == "fee_no_currency_sign" for f in a.flags)
    assert a.total.time.clock_minutes_max == 12
    assert a.total.fee.status == "not_stated"
    assert a.steps[0].person_responsible_raw == "Nurse"


def test_split_step_is_flagged_not_merged(tmp_path):
    _, b = split_workbook(make_health_sheet(tmp_path))
    assert [s.order for s in b.steps] == [1, 2, 3]
    assert b.steps[0].time.status == "not_stated" and b.steps[1].time.status == "ambiguous"
    assert b.steps[1].person_responsible_raw is None
    assert any(f["code"] == "possible_split_step" for f in b.flags)


def test_unmarked_extra_line_becomes_its_own_list_item_and_is_flagged(tmp_path):
    _, b = split_workbook(make_health_sheet(tmp_path))
    step = b.steps[2]
    assert step.agency_action_items == [
        "Examined",
        "Health history taken",
        "Report made and signed.",
    ]
    assert any(f["code"] == "orphan_row_glued" for f in b.flags)


def test_ambiguous_and_typo_durations_are_flagged(tmp_path):
    _, b = split_workbook(make_health_sheet(tmp_path))
    codes = [f["code"] for f in b.flags]
    assert "time_ambiguous" in codes
    assert "total_time_typo_normalized" in codes
    assert b.steps[1].time.raw == "1.15 min"
    assert b.steps[1].time.components == ()
    assert b.total.fee.items[0].amount_max == 30


# ------------------------------------------------------------------ real CHO workbook


def test_cho_ids_names_and_metadata(cho):
    assert [s.draft_id for s in cho] == [f"CHO-{i:02d}" for i in range(1, 16)]
    assert [s.source_number for s in cho] == list(range(1, 16))
    assert cho[0].name == "Routine Immunization"
    assert cho[3].name == "TB/HPN/Filariasis/Schstosomiasis/Leprosy treatment"
    assert cho[14].name == "Issuance of Medial Certificates for employment"  # source typo kept
    assert all(s.office == "City Health Office" for s in cho)
    assert all(s.classification == "SIMPLE" for s in cho)
    assert all(s.transaction_type == "Government to Citizen" for s in cho)
    assert cho[8].who_may_avail.startswith("This service caters to owners of business")
    assert cho[5].who_may_avail.startswith("This service caters to customers referred by")
    assert all(s.source_sheet == "HEALTH" for s in cho)


def test_cho_checklists_are_na_except_two(cho):
    with_reqs = {s.source_number: s for s in cho if s.requirements}
    assert set(with_reqs) == {10, 13}
    assert [(r.text, r.where_to_secure) for r in with_reqs[10].requirements] == [
        ("Application Form", None)
    ]
    assert [(r.text, r.where_to_secure) for r in with_reqs[13].requirements] == [
        ("Request to conduct a Post-Mortem Examination", "Philippine National Police")
    ]
    assert sum(any(f["code"] == "no_requirements_listed" for f in s.flags) for s in cho) == 13


def test_cho_step_counts(cho):
    assert [len(s.steps) for s in cho] == [3, 3, 2, 2, 3, 4, 3, 3, 3, 3, 1, 4, 3, 3, 3]


def test_cho_bullets(cho):
    assert cho[0].steps[0].agency_action_items == [
        "Client is taken in, mother baby book prepared and old record retrieved",
        "Client details entered in recording logbook",
        "Client weighed and vital signs taken and recorded in baby booklet",
    ]
    assert cho[0].steps[2].agency_action_items is None  # "Supplies are logged ..." has no bullets
    assert cho[14].steps[0].agency_action_items == [
        "Clients name is taken, forms are filled-in and vital signs taken",
        "Purpose for certification is taken",
        "Billing is given to the patient for payment at the treasurer’s office",
        "Request for diagnostic exams required given",
    ]
    # tab after the bullet (family planning) still splits cleanly
    assert cho[2].steps[0].agency_action_items[0].startswith("Client is taken in, old record")


def test_cho_fee_forms(cho):
    assert all(st.fees.status == "none" for s in cho[:10] for st in s.steps if st.fees.raw)
    dental = cho[10].steps[0].fees
    assert [(i.amount_max, i.note) for i in dental.items] == [(250, "no currency sign in source")]
    medical = cho[14]
    assert medical.steps[0].fees.items[0].amount_max == 30
    assert medical.total.fee.items[0].amount_max == 30
    assert medical.total.fee_text == "₱30.00"


def test_cho_duration_forms_and_flags(cho):
    assert cho[0].steps[0].time.components[0].unit == "minute"
    assert cho[0].steps[2].time.raw == "2 mins"
    assert cho[3].steps[1].time.status == "ambiguous"  # "1.15 min"
    assert cho[3].steps[1].time.raw == "1.15 min"
    assert any(f["code"] == "time_ambiguous" for f in cho[3].flags)
    assert cho[3].total.time.clock_minutes_max == 90  # "1 hour 30 min"
    d = cho[5].steps[2].time  # "2-3 days"
    assert (d.days_min, d.days_max) == (2, 3)
    assert cho[5].total.time.raw == "3 dsys"
    assert cho[5].total.time.normalizations == ("dsys -> day",)
    assert any(f["code"] == "total_time_typo_normalized" for f in cho[5].flags)
    assert cho[9].steps[1].time.days_max == 3


def test_cho_split_step_and_continuation_line(cho):
    lab = cho[5]  # Laboratory Services: row 73 has the label and person, row 74 the fee and time
    assert lab.steps[0].time.status == "not_stated" and lab.steps[0].fees.status == "not_stated"
    assert lab.steps[1].time.raw == "5 min" and lab.steps[1].person_responsible_raw is None
    assert any(f["code"] == "possible_split_step" for f in lab.flags)
    medico = cho[13]  # "Report made and signed." sits on its own row under step 2
    assert medico.steps[1].agency_action_items[-1] == "Report made and signed."
    assert any(f["code"] == "orphan_row_glued" for f in medico.flags)


def test_cho_source_row_bookkeeping(cho):
    assert cho[0].source_rows == (2, 13)
    assert cho[14].source_rows == (184, 195)
    assert cho[6].source_rows == (79, 91)  # pharmacy has a blank row after the title


# ------------------------------------------------------------------ real CSWDO workbook


def test_cswdo_metadata(cswdo):
    (svc,) = cswdo
    assert svc.draft_id == "CSWDO-01"
    assert svc.title == "1. Referrals" and svc.name == "Referrals"
    assert svc.office == "City Social Welfare Development Office"
    assert svc.classification == "COMPLEX"
    assert svc.transaction_type == "Government of Citizen"  # source wording kept
    assert svc.who_may_avail == "General Public"
    assert svc.description.startswith("This includes referrals to government hospitals")
    assert "necessary intervention/asssitance" in svc.description


def test_cswdo_wrapped_requirements_are_joined_and_flagged(cswdo):
    (svc,) = cswdo
    assert [(r.key, r.text, r.where_to_secure) for r in svc.requirements] == [
        ("1", "Brgy Certification as to residence", "Brgy Hall"),
        (
            "2",
            "Certification from Assessor's Office that client does not own real property",
            "City Assessor's Office",
        ),
        (
            "3",
            "Certification from BPLO that the client has no existing business",
            "BPLO",
        ),
        ("4", "Medical Abstract", None),
        ("5", "Death Certificate", None),
    ]
    assert sum(f["code"] == "requirement_fragment_glued" for f in svc.flags) == 2


def test_cswdo_wrapped_steps_become_four_steps(cswdo):
    (svc,) = cswdo
    assert len(svc.steps) == 4
    s1, s2, s3, s4 = svc.steps
    assert s1.citizen_action == "Present required document"
    assert s1.agency_action == (
        "Intake interview necessary information regarding the client and family are adduced"
    )
    assert s1.person_responsible_raw == "Personel in charge / Emergency Welfare Program implementer"
    assert s1.time.components[0].value_max == 30 and s1.time.raw == "30 min/s"
    assert s2.citizen_action == (
        "Clients referred to LTO for transportation assisstance / PCSO for financial (Medical) "
        "assistance / Missionaries of Charity for temporary placement / SOS for long term "
        "residential care"
    )
    assert s2.agency_action == (
        "Preparation of Social Case Study Report / Certificate of Indigency"
    )
    assert s2.time.components[0].unit == "week"
    assert s3.agency_action == (
        "Release of needed documents (SCSR / Certificate of Indigency / Referral Letter)"
    )
    assert s3.time.raw == "10 min/s"
    assert s4.agency_action == "Monitoring" and s4.citizen_action is None
    assert s4.time.raw == "1 hour (Once a month)"
    assert [s.citizen_step_group for s in svc.steps] == [1, 2, 3, 3]
    assert sum(f["code"] == "orphan_row_glued" for f in svc.flags) == 13


def test_cswdo_total(cswdo):
    (svc,) = cswdo
    assert svc.total.time_text == "1 week, 1 hour, 40 minutes"
    assert svc.total.time.weeks_max == 1 and svc.total.time.clock_minutes_max == 100
    assert svc.total.fee.status == "not_stated"
    assert svc.source_rows == (3, 38)
