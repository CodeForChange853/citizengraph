"""Splitter: data/raw workbooks -> rough per-service drafts (structure only, no curation)."""

from pathlib import Path

import openpyxl
import pytest
import yaml

from citizengraph.parsing.splitter import ServiceDraft, split_workbook, write_drafts

RAW = Path(__file__).resolve().parents[1] / "data" / "raw"


@pytest.fixture(scope="module")
def bplo() -> list[ServiceDraft]:
    return split_workbook(RAW / "BPLO-CC.xlsx")


@pytest.fixture(scope="module")
def lcro() -> list[ServiceDraft]:
    return split_workbook(RAW / "LCRO-CC.xlsx")


def make_sheet(tmp_path: Path) -> Path:
    """A miniature service block that exercises merged cells and every header row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"] = "TEST OFFICE"
    ws["A2"] = "List of Services"
    ws["A3"] = "1. Sample Service"
    ws["A5"] = "1. Sample Service"
    ws["A6"] = "A sample description."
    ws["A7"], ws["B7"] = "Name of frontline service:", "SAMPLE SERVICE"
    ws["A8"], ws["B8"] = "Office or Department:", "TEST OFFICE"
    ws["A9"], ws["B9"] = " Classification:", "Simple"
    ws["A10"], ws["B10"] = "Type of Transaction:", "G2C"
    ws["A11"] = "Who may avail:"
    ws["A12"], ws["C12"] = "CHECKLIST OF REQUIREMENTS", "WHERE TO SECURE"
    ws["A13"], ws["C13"] = "A. First doc", "Office X"
    ws["A14"], ws["C14"] = "B. Second doc", None
    ws["A15"], ws["C15"] = "6.1. Sub doc", "3. Office Y"
    ws["A16"], ws["C16"] = "Third doc, unlabelled", "- - -"
    ws["A17"], ws["C17"] = "Group heading:", None
    ws["A18"], ws["C18"] = "     •    Bullet doc", "Office Z"
    ws.merge_cells("C13:C14")
    for col, head in zip(
        "ABCDE",
        [
            "CITIZEN STEPS",
            "AGENCY ACTION",
            "FEES TO BE PAID",
            "PROCESSING TIME",
            "PERSON RESPONSIBLE",
        ],
    ):
        ws[f"{col}20"] = head
    ws["A21"], ws["B21"], ws["C21"], ws["D21"], ws["E21"] = (
        "A. Submit",
        "A.1. Check",
        "₱10.00\n*ALL FEES MUST BE SETTLED AT THE CTO*",
        "5 minutes",
        "Juan Dela Cruz",
    )
    ws["B22"], ws["D22"], ws["E22"] = "A.2. Encode", "2 minutes", "Any staff"
    ws["A23"], ws["B23"], ws["C23"], ws["D23"], ws["E23"] = (
        "B. Receive",
        "B.1. Print",
        "- - -",
        "3 minutes",
        "Registrar",
    )
    ws["B24"], ws["B25"] = "B.2. Sign", "B.3. Release"
    ws["D24"] = "1 hour"
    ws["E24"] = "Registrar"
    ws.merge_cells("D24:D25")
    ws["A26"], ws["C26"], ws["D26"] = (
        "TOTAL:",
        "₱10.00",
        " 1 hour, 10 minutes (under normal conditions)",
    )
    path = tmp_path / "TEST-CC.xlsx"
    wb.save(path)
    return path


def test_synthetic_block_metadata_and_totals(tmp_path):
    (svc,) = split_workbook(make_sheet(tmp_path))
    assert svc.draft_id == "TEST-01"
    assert svc.source_number == 1
    assert svc.title == "1. Sample Service"
    assert svc.description == "A sample description."
    assert svc.name == "SAMPLE SERVICE"
    assert svc.office == "TEST OFFICE"
    assert svc.classification == "SIMPLE"
    assert svc.classification_raw == "Simple"
    assert svc.transaction_type == "G2C"
    assert svc.who_may_avail is None
    assert svc.total.fee.items[0].amount_max == 10
    assert svc.total.time.components[0].unit == "hour"
    assert svc.total.time_text == "1 hour, 10 minutes (under normal conditions)"


def test_synthetic_requirements(tmp_path):
    (svc,) = split_workbook(make_sheet(tmp_path))
    reqs = svc.requirements
    assert [r.key for r in reqs] == ["A", "B", "6.1", None, None, None]
    assert reqs[0].text == "First doc" and reqs[0].where_to_secure == "Office X"
    # merged "where to secure" is shared with the row above, flagged, not silently copied
    assert reqs[1].where_to_secure == "Office X"
    assert reqs[1].where_shared_with == reqs[0].ref
    # list markers are stripped from "where to secure"; "- - -" means null
    assert reqs[2].where_to_secure == "Office Y"
    assert reqs[3].where_to_secure is None
    assert reqs[4].kind == "group" and reqs[5].kind == "bullet"
    assert reqs[5].parent_ref == reqs[4].ref
    assert reqs[5].text == "Bullet doc"


def test_synthetic_steps_merged_cells_are_counted_once(tmp_path):
    (svc,) = split_workbook(make_sheet(tmp_path))
    steps = svc.steps
    assert [s.order for s in steps] == [1, 2, 3, 4, 5]
    assert [s.label for s in steps] == ["A.1.", "A.2.", "B.1.", "B.2.", "B.3."]
    assert [s.citizen_step_group for s in steps] == [1, 1, 2, 2, 2]
    assert steps[0].citizen_action == "A. Submit" and steps[1].citizen_action is None
    assert steps[0].fees.items[0].amount_max == 10
    assert steps[0].fees.notes == ("ALL FEES MUST BE SETTLED AT THE CTO",)
    assert steps[3].time.components[0].unit == "hour"
    assert steps[4].time.status == "not_stated"  # second row of a merged D24:D25
    assert steps[4].time_shared_from == 4
    assert steps[3].time_span == 2
    # person responsible kept raw (never citizen-facing) and named for the draft only
    assert steps[0].person_responsible_raw == "Juan Dela Cruz"
    assert steps[4].person_responsible_raw is None


def test_ids_come_from_file_name_and_service_number(bplo, lcro):
    assert [s.draft_id for s in bplo] == [f"BPLO-{i:02d}" for i in range(1, 8)]
    assert [s.draft_id for s in lcro] == [f"LCRO-{i:02d}" for i in range(1, 18)]


def test_all_24_services_found(bplo, lcro):
    assert len(bplo) == 7 and len(lcro) == 17


def test_bplo_metadata(bplo):
    assert bplo[0].name == "Business Permit"
    assert bplo[0].office.startswith("Business Permits & Licensing Office")
    assert bplo[0].who_may_avail == "Business Owners (New & Renewals)"
    assert all(s.classification == "SIMPLE" for s in bplo)
    # charter_data.md section 3: who may avail is blank for BPLO 3 to 7 (Cockfight has text)
    assert [bplo[i].who_may_avail for i in (2, 3, 5, 6)] == [None, None, None, None]
    assert bplo[4].who_may_avail == "Open to all cockfight enthusiasts"


def test_lcro_classification_matches_charter_data(lcro):
    complex_numbers = {s.source_number for s in lcro if s.classification == "COMPLEX"}
    assert complex_numbers == {2, 5, 7, 15, 16}


def test_stated_totals_match_charter_data_table(lcro):
    expected = {  # docs/charter_data.md section 2: fee max, (days_max, clock minutes max)
        1: (None, (0, 38)),
        2: (None, (10, 33)),
        3: (300, (0, 66)),
        4: (None, (0, 31)),
        5: (None, (10, 28)),
        6: (None, (0, 38)),
        7: (None, (10, 38)),
        8: (200, (0, 36)),
        9: (110, (0, 28)),
        10: (300, (0, 22 * 60 + 51)),
        11: (220, (0, 74)),
        12: (580, (0, 59)),
        13: (540, (0, 74)),
        14: (300, (0, 74)),
        15: (3000, (10, 114)),
        16: (None, (0, 59)),
        17: (295, (10, 26)),
    }
    for svc in lcro:
        fee_max, (days, minutes) = expected[svc.source_number]
        items = svc.total.fee.items
        if fee_max is None:
            assert svc.total.fee.status == "none", svc.draft_id
        else:
            assert max(i.amount_max for i in items) == fee_max, svc.draft_id
        assert svc.total.time.days_max == days, svc.draft_id
        assert svc.total.time.clock_minutes_max == minutes, svc.draft_id


def test_bplo_stated_totals(bplo):
    assert bplo[0].total.fee.items[0].amount_max == 235.5
    assert bplo[0].total.time.clock_minutes_max == 37
    assert [(i.amount_max, i.qualifier) for i in bplo[1].total.fee.items] == [
        (215, "company"),
        (310, "individual"),
    ]
    assert bplo[5].total.time.clock_minutes_max == 16
    assert bplo[6].total.fee.status == "not_stated"
    assert bplo[6].total.time.clock_minutes_max == 25


def test_step_counts_spot_check(bplo, lcro):
    assert len(bplo[0].steps) == 10
    assert len(lcro[0].steps) == 8
    assert len(lcro[2].steps) == 9
    assert len(lcro[10].steps) == 11
    assert len(lcro[16].steps) == 9


def test_lcro_marriage_license_details(lcro):
    svc = lcro[2]
    assert svc.name == "APPLICATION FOR MARRIAGE LICENSE"
    assert svc.requirements[1].text == "Birth Certificate"
    assert [r.key for r in svc.requirements if r.key and r.key.startswith("6.")] == [
        "6.1",
        "6.2",
        "6.3",
    ]
    six = next(r for r in svc.requirements if r.key == "6")
    assert all(r.parent_ref == six.ref for r in svc.requirements if r.key in ("6.1", "6.2", "6.3"))
    posting = svc.steps[6]  # 5.3 Posting of Notice of Marriage: no time stated
    assert posting.agency_action.startswith("Posting of Notice of Marriage")
    assert posting.time.status == "not_stated"
    assert svc.steps[5].person_responsible_raw is None  # 5.2 has no responsible person


def test_merged_time_is_attributed_once(bplo, lcro):
    occ = bplo[1]  # D70:D72 "2 minutes" shared by E.1, E.2, E.3
    assert [s.time.status for s in occ.steps[-3:]] == ["stated", "not_stated", "not_stated"]
    assert occ.steps[-3].time_span == 3
    assert occ.steps[-2].time_shared_from == occ.steps[-3].order
    breqs = lcro[16]
    assert sum(1 for s in breqs.steps if s.time.status == "stated" and s.time_span == 4) == 1


def test_lcro_court_order_fee_list_kept_as_text(lcro):
    step = next(s for s in lcro[10].steps if "local tax code" in (s.agency_action or ""))
    assert "Annulment" in step.agency_action and "Presumptive Death" in step.agency_action
    assert step.fees.status == "not_stated"


def test_flat_fees_list_and_yaml_round_trip(bplo, tmp_path):
    svc = bplo[0]
    assert [f["label"] for f in svc.to_dict()["fees"]] == [
        "Zoning",
        "Sanitary Services",
        "Business Name Clearance",
        "Signboard",
    ]
    paths = write_drafts(bplo, tmp_path)
    assert len(paths) == 7
    loaded = yaml.safe_load(paths[0].read_text(encoding="utf-8"))
    assert loaded["draft_id"] == "BPLO-01"
    assert loaded["name"] == "Business Permit"
    assert loaded["total"]["time"]["day_type"] == "unknown"
    assert loaded["steps"][0]["time"]["components"] == [
        {"value_min": 2, "value_max": 2, "unit": "minute"}
    ]
    assert paths[0].read_text(encoding="utf-8").startswith("# DRAFT")
    assert write_drafts(bplo, tmp_path) == paths  # deterministic, no timestamps
